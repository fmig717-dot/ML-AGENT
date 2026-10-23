"""
Инструменты файловой системы для Claude: чтение файлов skill'ов (только чтение,
они зашиты в образ контейнера) и чтение/запись рабочих файлов на постоянном
диске DATA_DIR (buyers.xlsx, kp_state.json, вложения, заметки).

Везде защита от выхода за пределы разрешённой директории (path traversal).
"""
from __future__ import annotations

import os

from app.config import DATA_DIR, SKILLS_DIR

TEXT_EXTENSIONS = {".md", ".json", ".txt", ".py", ".yml", ".yaml", ".csv"}
MAX_READ_BYTES = 200_000  # защита от случайного чтения огромного/бинарного файла


def _resolve(base: str, rel_path: str) -> str:
    base = os.path.abspath(base)
    full = os.path.abspath(os.path.join(base, rel_path))
    if not (full == base or full.startswith(base + os.sep)):
        raise ValueError(f"Путь '{rel_path}' выходит за пределы разрешённой директории")
    return full


def _tree(base: str, rel_path: str = "") -> list[str]:
    root = _resolve(base, rel_path)
    out = []
    if not os.path.exists(root):
        return out
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for fn in sorted(filenames):
            if fn.startswith("."):
                continue
            full = os.path.join(dirpath, fn)
            out.append(os.path.relpath(full, base))
    return sorted(out)


def list_skill_tree(sub_path: str = "") -> dict:
    """Список всех файлов внутри .claude/skills/<sub_path> (рекурсивно)."""
    return {"base": "skills", "files": _tree(SKILLS_DIR, sub_path)}


def read_skill_file(path: str) -> dict:
    """Читает текстовый файл skill'а (SKILL.md, rules.md, шаблон, скрипт) целиком."""
    full = _resolve(SKILLS_DIR, path)
    if not os.path.isfile(full):
        return {"error": f"Файл не найден: {path}"}
    size = os.path.getsize(full)
    if size > MAX_READ_BYTES:
        return {"error": f"Файл слишком большой ({size} байт) для чтения целиком"}
    with open(full, encoding="utf-8", errors="replace") as f:
        return {"path": path, "content": f.read()}


def list_data_tree(sub_path: str = "") -> dict:
    """Список файлов в рабочей директории DATA_DIR (buyers.xlsx, kp_state.json и т.п.)."""
    return {"base": "data", "files": _tree(DATA_DIR, sub_path)}


def read_data_file(path: str) -> dict:
    """
    Читает ТЕКСТОВЫЙ файл из рабочей директории (.json/.md/.txt/.csv).
    Для .xlsx используй read_xlsx_sheet — бинарные файлы этим инструментом не читаются.
    """
    full = _resolve(DATA_DIR, path)
    ext = os.path.splitext(full)[1].lower()
    if ext not in TEXT_EXTENSIONS:
        return {"error": f"'{ext}' — не текстовый формат. Для .xlsx используй read_xlsx_sheet."}
    if not os.path.isfile(full):
        return {"error": f"Файл не найден: {path}"}
    size = os.path.getsize(full)
    if size > MAX_READ_BYTES:
        return {"error": f"Файл слишком большой ({size} байт) для чтения целиком"}
    with open(full, encoding="utf-8", errors="replace") as f:
        return {"path": path, "content": f.read()}


def write_data_file(path: str, content: str) -> dict:
    """Создаёт или перезаписывает текстовый файл (.json/.md/.txt) в рабочей директории."""
    full = _resolve(DATA_DIR, path)
    ext = os.path.splitext(full)[1].lower()
    if ext not in TEXT_EXTENSIONS:
        return {"error": f"'{ext}' — запись разрешена только для текстовых форматов"}
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "w", encoding="utf-8") as f:
        f.write(content)
    return {"path": path, "bytes_written": len(content.encode("utf-8"))}


MAX_DOC_READ_BYTES = 20_000_000  # документы (спецификации/КП) могут быть крупнее текстовых файлов


def read_pdf_file(path: str, max_pages: int = 50) -> dict:
    """
    Извлекает текст из PDF (спецификация, анализ рынка, готовое КП и т.п.), присланного
    пользователем через Telegram и сохранённого в рабочей директории (см. list_data_tree
    по 'inbox/<chat_id>'). Не извлекает изображения/таблицы как структуру — только
    текстовый слой; если PDF — это скан без текстового слоя, вернётся пустой/короткий
    текст, об этом нужно честно сообщить пользователю, а не выдумывать содержание.
    """
    from pypdf import PdfReader

    full = _resolve(DATA_DIR, path)
    if not os.path.isfile(full):
        return {"error": f"Файл не найден: {path}"}
    size = os.path.getsize(full)
    if size > MAX_DOC_READ_BYTES:
        return {"error": f"Файл слишком большой ({size} байт) для чтения целиком"}
    try:
        reader = PdfReader(full)
    except Exception as exc:
        return {"error": f"Не удалось открыть PDF: {type(exc).__name__}: {exc}"}

    total_pages = len(reader.pages)
    texts = []
    for i, page in enumerate(reader.pages):
        if i >= max_pages:
            texts.append(f"... (обрезано на странице {max_pages} из {total_pages})")
            break
        texts.append(page.extract_text() or "")
    return {"path": path, "total_pages": total_pages, "content": "\n\n".join(texts)}


def read_docx_file(path: str) -> dict:
    """Извлекает текст (параграфы и таблицы) из .docx, присланного пользователем."""
    import docx

    full = _resolve(DATA_DIR, path)
    if not os.path.isfile(full):
        return {"error": f"Файл не найден: {path}"}
    size = os.path.getsize(full)
    if size > MAX_DOC_READ_BYTES:
        return {"error": f"Файл слишком большой ({size} байт) для чтения целиком"}
    try:
        document = docx.Document(full)
    except Exception as exc:
        return {"error": f"Не удалось открыть DOCX: {type(exc).__name__}: {exc}"}

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return {"path": path, "content": "\n".join(parts)}


def read_xlsx_sheet(path: str, sheet_name: str | None = None, max_rows: int = 300) -> dict:
    """
    Читает .xlsx из рабочей директории. Без sheet_name — возвращает список листов.
    С sheet_name — возвращает строки листа (включая заголовок) как список списков.
    """
    import openpyxl

    full = _resolve(DATA_DIR, path)
    if not os.path.isfile(full):
        return {"error": f"Файл не найден: {path}"}
    wb = openpyxl.load_workbook(full, data_only=True)
    if sheet_name is None:
        return {"sheets": wb.sheetnames}
    if sheet_name not in wb.sheetnames:
        return {"error": f"Листа '{sheet_name}' нет. Доступные листы: {wb.sheetnames}"}
    ws = wb[sheet_name]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            rows.append(["... (обрезано, увеличьте max_rows или читайте по частям)"])
            break
        rows.append(["" if c is None else c for c in row])
    return {"sheet": sheet_name, "row_count_returned": len(rows), "rows": rows}
