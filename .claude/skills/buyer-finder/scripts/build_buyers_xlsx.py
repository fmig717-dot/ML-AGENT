"""
Строит buyers.xlsx из buyers.json (карточек покупателей).

Формат вывода — книга из шести листов:
  1. "Buyers"            — основная таблица покупателей (фильтр по стране/региону,
                           статус переговоров и QC-статусы через выпадающие списки).
  2. "История контактов" — лог реальных касаний с каждой компанией.
  3. "Лог запусков"      — аудит запусков агента поиска (buyer-finder).
  4. "Рынок и тренды"    — лог рыночных исследований агента-надзирателя
                           (market-supervisor): тренды, цены, источники.
  5. "QC лог"             — детальные записи проверок качества (данных о покупателе
                           и текста писем) от агента-надзирателя.
  6. "KPI и токены"       — произвольные метрики проекта (найдено/принято/отправлено,
                           расход токенов и т.п.), которые ведёт агент-надзиратель.

Листы 2-6 заполняются другими скриптами/агентами напрямую через openpyxl между
пересборками — этот скрипт их не генерирует, только создаёт (если их ещё нет) и
ВСЕГДА переносит их прежнее содержимое при пересборке, чтобы ничего не терялось.

Важно: скрипт запускаемый повторно НЕ должен затирать то, что было вписано в
buyers.xlsx руками или другими агентами (статус переговоров, QC-статусы, номер КП,
волна, и все строки листов 2-6). Поэтому перед перезаписью файла скрипт читает
существующий buyers.xlsx (если он есть) и переносит эти данные в новую книгу:
Excel-файл считается источником правды для мутируемых полей, buyers.json — только
для состава компаний и их карточек (название/контакты/заметки/fit_reasoning).

Использование:
    python3 build_buyers_xlsx.py <путь_к_buyers.json> <путь_к_buyers.xlsx> [run_log.json]

run_log.json (опционально) — см. пример в SKILL.md buyer-finder. Если файл не передан —
новая строка в "Лог запусков" не добавляется, но все предыдущие записи сохраняются.

Если аргументы не переданы, по умолчанию ищет buyers.json и пишет buyers.xlsx
рядом (в текущей рабочей папке, в output/).
"""
import json
import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEFAULT_INPUT = os.path.join("output", "buyers.json")
DEFAULT_OUTPUT = os.path.join("output", "buyers.xlsx")

input_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_INPUT
output_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT
run_log_path = sys.argv[3] if len(sys.argv) > 3 else None

STATUS_OPTIONS = [
    "Не контактировали",
    "Драфт КП создан (не отправлен)",
    "Soft offer отправлен",
    "Ждём ответа",
    "В переговорах",
    "Запрошен образец/тестовая партия",
    "Образец отправлен",
    "Согласование контракта",
    "Сделка закрыта",
    "Отказ/не подходит",
    "Ошибка - пропущено (авто)",
]

QC_OPTIONS = [
    "",
    "На проверке",
    "Принято",
    "Нужна доработка",
    "Отклонено",
]

# Поля, которые пишут не buyer-finder, а другие агенты (offer-drafter, market-supervisor)
# напрямую в buyers.xlsx — при пересборке книги они всегда ценнее одноимённых
# значений из buyers.json, иначе пересборка списка покупателей сносит их прогресс.
MUTABLE_FIELDS = ("status", "last_contact_date", "next_step", "kp_number", "wave", "qc_data", "qc_email")

BUYERS_HEADERS = [
    "Компания", "Продукт", "Страна", "Регион/город", "Сайт", "Email", "Телефон",
    "Контактное лицо", "Почему подходит", "Заметки для письма",
    "Статус переговоров", "Дата последнего контакта", "Следующий шаг",
    "Номер КП", "Волна", "QC данных", "QC письма",
]
BUYERS_FIELD_ORDER = [
    "company", "product", "country", "region", "website", "email", "phone",
    "contact_person", "fit_reasoning", "notes",
    "status", "last_contact_date", "next_step",
    "kp_number", "wave", "qc_data", "qc_email",
]
BUYERS_WIDTHS = {
    "A": 32, "B": 14, "C": 16, "D": 22, "E": 26, "F": 26, "G": 22, "H": 20,
    "I": 46, "J": 46, "K": 26, "L": 18, "M": 34, "N": 18, "O": 10, "P": 16, "Q": 16,
}

# Листы 2-6: заголовки и ширины колонок. Эти листы скрипт только создаёт и сохраняет —
# новые строки в них (кроме "Лог запусков", см. ниже) добавляют другие скрипты/агенты.
EXTRA_SHEETS = {
    "История контактов": {
        "headers": ["Дата", "Компания", "Канал", "Содержание", "Автор", "Следующий шаг"],
        "widths": {"A": 14, "B": 32, "C": 16, "D": 60, "E": 18, "F": 34},
    },
    "Лог запусков": {
        "headers": ["Дата запуска", "Продукт", "Критерии/регион поиска", "Найдено", "Принято", "Отклонено (с причиной)", "Источники"],
        "widths": {"A": 14, "B": 34, "C": 34, "D": 10, "E": 10, "F": 50, "G": 34},
    },
    "Рынок и тренды": {
        "headers": ["Дата", "Продукт", "Источник", "Инсайт/тренд", "Рекомендация (цена/подход)", "Статус"],
        "widths": {"A": 14, "B": 14, "C": 34, "D": 50, "E": 50, "F": 26},
    },
    "QC лог": {
        "headers": ["Дата", "Компания", "Волна", "Тип проверки", "Статус", "Комментарий", "Проверил"],
        "widths": {"A": 14, "B": 32, "C": 10, "D": 20, "E": 18, "F": 50, "G": 18},
    },
    "KPI и токены": {
        "headers": ["Дата", "Волна", "Метрика", "Значение", "Комментарий"],
        "widths": {"A": 14, "B": 10, "C": 30, "D": 18, "E": 50},
    },
}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")


def read_existing_state(path):
    """Читает старый buyers.xlsx (если есть) и возвращает то, что нужно сохранить:
    - per_company: {company_name: {status, last_contact_date, next_step, kp_number, wave, qc_data, qc_email}}
    - extra_rows: {sheet_name: [строки без заголовка]} для каждого листа из EXTRA_SHEETS
    """
    per_company = {}
    extra_rows = {name: [] for name in EXTRA_SHEETS}

    if not path or not os.path.exists(path):
        return per_company, extra_rows

    try:
        wb_old = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"Предупреждение: не удалось открыть старый {path} ({exc}), продолжаю без переноса данных.")
        return per_company, extra_rows

    if "Buyers" in wb_old.sheetnames:
        ws_old = wb_old["Buyers"]
        header_row = [c.value for c in ws_old[1]]
        header_to_field = {
            "Статус переговоров": "status",
            "Дата последнего контакта": "last_contact_date",
            "Следующий шаг": "next_step",
            "Номер КП": "kp_number",
            "Волна": "wave",
            "QC данных": "qc_data",
            "QC письма": "qc_email",
        }
        try:
            idx_company = header_row.index("Компания")
        except ValueError:
            idx_company = None
        idx_by_field = {}
        for header_title, field in header_to_field.items():
            if header_title in header_row:
                idx_by_field[field] = header_row.index(header_title)

        if idx_company is not None:
            for row in ws_old.iter_rows(min_row=2, values_only=True):
                if row[idx_company] is None or str(row[idx_company]).strip() == "":
                    continue
                per_company[str(row[idx_company]).strip()] = {
                    field: (row[idx] if row[idx] else "")
                    for field, idx in idx_by_field.items()
                }

    for sheet_name in EXTRA_SHEETS:
        if sheet_name in wb_old.sheetnames:
            ws_extra = wb_old[sheet_name]
            for row in ws_extra.iter_rows(min_row=2, values_only=True):
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                extra_rows[sheet_name].append(row)

    return per_company, extra_rows


def style_header(ws, headers):
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def write_extra_sheet(wb, sheet_name, rows):
    config = EXTRA_SHEETS[sheet_name]
    ws = wb.create_sheet(sheet_name)
    style_header(ws, config["headers"])
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row[: len(config["headers"])], start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP_LEFT
    for col, width in config["widths"].items():
        ws.column_dimensions[col].width = width
    last_row = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(config['headers']))}{last_row}"
    return ws


def main():
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)

    preserved_status, extra_rows = read_existing_state(output_path)

    wb = openpyxl.Workbook()

    # --- Лист 1: Buyers ---
    ws = wb.active
    ws.title = "Buyers"
    style_header(ws, BUYERS_HEADERS)

    for row_idx, entry in enumerate(data, start=2):
        company_name = str(entry.get("company", "")).strip()
        preserved = preserved_status.get(company_name)

        for col_idx, field in enumerate(BUYERS_FIELD_ORDER, start=1):
            if preserved and field in MUTABLE_FIELDS and preserved.get(field):
                value = preserved.get(field)
            else:
                value = entry.get(field, "")
                if field == "status" and not value:
                    value = STATUS_OPTIONS[0]

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP_LEFT

        website = entry.get("website", "") or ""
        email = entry.get("email", "") or ""
        website_cell = ws.cell(row=row_idx, column=BUYERS_FIELD_ORDER.index("website") + 1)
        email_cell = ws.cell(row=row_idx, column=BUYERS_FIELD_ORDER.index("email") + 1)
        if website.startswith("http"):
            website_cell.hyperlink = website
            website_cell.font = LINK_FONT
        if "@" in email and "не найдено" not in email:
            email_cell.hyperlink = f"mailto:{email}"
            email_cell.font = LINK_FONT

    for col, width in BUYERS_WIDTHS.items():
        ws.column_dimensions[col].width = width
    for row_idx in range(2, len(data) + 2):
        ws.row_dimensions[row_idx].height = 90

    last_col_letter = get_column_letter(len(BUYERS_HEADERS))
    last_row = max(len(data) + 1, 2)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    def add_dropdown(field_name, options, error_label):
        col_letter = get_column_letter(BUYERS_FIELD_ORDER.index(field_name) + 1)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(o for o in options if o) + '"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.error = error_label
        dv.errorTitle = "Недопустимое значение"
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}500")

    add_dropdown("status", STATUS_OPTIONS, "Выберите статус из списка")
    add_dropdown("qc_data", QC_OPTIONS, "Выберите QC-статус из списка")
    add_dropdown("qc_email", QC_OPTIONS, "Выберите QC-статус из списка")

    # --- Листы 2-6 ---
    for sheet_name in EXTRA_SHEETS:
        rows = list(extra_rows[sheet_name])
        if sheet_name == "Лог запусков" and run_log_path and os.path.exists(run_log_path):
            with open(run_log_path, encoding="utf-8") as f:
                run_log = json.load(f)
            rows.append((
                run_log.get("date", ""),
                run_log.get("product", ""),
                run_log.get("criteria_region", ""),
                run_log.get("found", ""),
                run_log.get("accepted", ""),
                run_log.get("rejected_reasons", ""),
                run_log.get("sources", ""),
            ))
        write_extra_sheet(wb, sheet_name, rows)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    counts = ", ".join(f"{name}={len(extra_rows[name])}" for name in EXTRA_SHEETS)
    print(f"Saved {output_path}: {len(data)} покупателей; сохранённые строки по листам: {counts}.")


if __name__ == "__main__":
    main()
