"""
Записывает результат создания ОДНОГО драфта письма в buyers.xlsx:
  - обновляет на листе "Buyers" строку компании: Статус переговоров, Номер КП, Волну,
    Следующий шаг (и, если передана, Дату последнего контакта);
  - добавляет строку в лист "История контактов".

Вызывай это СРАЗУ после успешного create_draft по одной компании — чтобы прогресс
волны фиксировался по одному письму, а не пачкой в конце (если агент упадёт/остановится
по правилу 6 на середине волны, уже обработанные компании должны быть отмечены).

Использование:
    python3 record_draft.py <buyers.xlsx> \
        --company "Degam Oil and Chemical Industry (Degam Yağ ve Kimya Sanayi)" \
        --kp "CFA/2610-22_01" --wave 1 \
        --status "Драфт КП создан (не отправлен)" \
        --next-step "Проверить и отправить вручную" \
        --history-channel "Email (черновик)" \
        --history-note "Создан драфт КП CFA/2610-22_01, не отправлен — ожидает проверки" \
        [--author "Agent"] [--last-contact-date ""] [--history-date "22.10.2026"]

Если компания с таким названием не найдена на листе Buyers — скрипт завершается с
ошибкой (exit code 1) и ничего не меняет. Сверяй точное название из той же строки
таблицы, которую показывал пользователю при подтверждении волны.
"""
import argparse
import os
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BODY_FONT = Font(name="Arial", size=10)
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path")
    parser.add_argument("--company", required=True)
    parser.add_argument("--kp", required=True, dest="kp_number")
    parser.add_argument("--wave", required=True, type=int)
    parser.add_argument("--status", required=True)
    parser.add_argument("--next-step", default="")
    parser.add_argument("--last-contact-date", default=None)
    parser.add_argument("--history-channel", default="Email (черновик)")
    parser.add_argument("--history-note", required=True)
    parser.add_argument("--author", default="Agent")
    parser.add_argument("--history-date", default=None)
    args = parser.parse_args()

    if not os.path.exists(args.xlsx_path):
        print(f"Ошибка: файл {args.xlsx_path} не найден.")
        sys.exit(1)

    wb = openpyxl.load_workbook(args.xlsx_path)

    if "Buyers" not in wb.sheetnames:
        print("Ошибка: на листе нет вкладки 'Buyers'.")
        sys.exit(1)

    ws = wb["Buyers"]
    header_row = [c.value for c in ws[1]]
    needed = ["Компания", "Статус переговоров", "Дата последнего контакта", "Следующий шаг", "Номер КП", "Волна"]
    missing = [h for h in needed if h not in header_row]
    if missing:
        print(f"Ошибка: на листе Buyers не найдены колонки {missing}. Пересобери buyers.xlsx свежим скриптом build_buyers_xlsx.py.")
        sys.exit(1)

    idx_company = header_row.index("Компания")
    idx_status = header_row.index("Статус переговоров")
    idx_date = header_row.index("Дата последнего контакта")
    idx_next = header_row.index("Следующий шаг")
    idx_kp = header_row.index("Номер КП")
    idx_wave = header_row.index("Волна")

    target_row = None
    for row in ws.iter_rows(min_row=2):
        cell_value = row[idx_company].value
        if cell_value is not None and str(cell_value).strip() == args.company.strip():
            target_row = row
            break

    if target_row is None:
        print(f"Ошибка: компания '{args.company}' не найдена на листе Buyers. Ничего не изменено.")
        sys.exit(1)

    target_row[idx_status].value = args.status
    target_row[idx_next].value = args.next_step
    target_row[idx_kp].value = args.kp_number
    target_row[idx_wave].value = args.wave
    if args.last_contact_date is not None:
        target_row[idx_date].value = args.last_contact_date

    # --- История контактов: добавить строку в конец ---
    if "История контактов" not in wb.sheetnames:
        print("Ошибка: на книге нет листа 'История контактов'. Пересобери buyers.xlsx свежим скриптом build_buyers_xlsx.py.")
        sys.exit(1)

    ws_hist = wb["История контактов"]
    next_hist_row = ws_hist.max_row + 1 if ws_hist.max_row and ws_hist.max_row >= 1 else 2
    # если лист пустой (только заголовок или вообще пусто) — пишем во вторую строку
    if ws_hist.max_row < 2:
        next_hist_row = 2

    history_date = args.history_date or date.today().strftime("%d.%m.%Y")
    values = [history_date, args.company, args.history_channel, args.history_note, args.author, args.next_step]
    for col_idx, value in enumerate(values, start=1):
        cell = ws_hist.cell(row=next_hist_row, column=col_idx, value=value)
        cell.font = BODY_FONT
        cell.alignment = WRAP_TOP_LEFT

    last_col_letter = get_column_letter(ws_hist.max_column)
    ws_hist.auto_filter.ref = f"A1:{last_col_letter}{max(next_hist_row, 2)}"

    wb.save(args.xlsx_path)
    print(f"OK: '{args.company}' -> статус='{args.status}', КП={args.kp_number}, волна={args.wave}; добавлена запись в Историю контактов (строка {next_hist_row}).")


if __name__ == "__main__":
    main()
