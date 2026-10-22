"""
Универсальный скрипт: добавляет одну строку в конец любого из листов
"Рынок и тренды" / "QC лог" / "KPI и токены" (или любого другого листа buyers.xlsx,
если у него есть строка заголовков в первой строке).

В отличие от build_buyers_xlsx.py (который пересобирает всю книгу из buyers.json),
этот скрипт просто дописывает одну строку в существующий файл — используй его для
логирования рыночных инсайтов, QC-вердиктов и KPI-метрик между пересборками.

Использование:
    python3 append_row.py <buyers.xlsx> "<Имя листа>" --field "Заголовок колонки=значение" [...]

Пример:
    python3 append_row.py output/buyers.xlsx "QC лог" \
      --field "Дата=22.10.2026" \
      --field "Компания=Degam Oil and Chemical Industry (Degam Yağ ve Kimya Sanayi)" \
      --field "Волна=1" \
      --field "Тип проверки=Текст письма" \
      --field "Статус=Принято" \
      --field "Комментарий=Подпись верна, вложения совпадают с manifest, дат в тексте нет" \
      --field "Проверил=market-supervisor"

Поля, которые есть в заголовке листа, но не переданы через --field, останутся пустыми
в новой строке. Поле, переданное через --field, но отсутствующее в заголовке листа —
ошибка (чтобы не плодить опечатки в названиях колонок).
"""
import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BODY_FONT = Font(name="Arial", size=10)
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path")
    parser.add_argument("sheet_name")
    parser.add_argument("--field", action="append", required=True, help="Заголовок=значение, можно повторять")
    args = parser.parse_args()

    if not os.path.exists(args.xlsx_path):
        print(f"Ошибка: файл {args.xlsx_path} не найден.")
        sys.exit(1)

    wb = openpyxl.load_workbook(args.xlsx_path)
    if args.sheet_name not in wb.sheetnames:
        print(f"Ошибка: лист '{args.sheet_name}' не найден в {args.xlsx_path}. Доступные листы: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[args.sheet_name]
    header_row = [c.value for c in ws[1]]

    field_values = {}
    for item in args.field:
        if "=" not in item:
            print(f"Ошибка: --field '{item}' должен быть в формате 'Заголовок=значение'.")
            sys.exit(1)
        key, value = item.split("=", 1)
        key = key.strip()
        if key not in header_row:
            print(f"Ошибка: колонка '{key}' не найдена в заголовке листа '{args.sheet_name}'. Заголовки: {header_row}")
            sys.exit(1)
        field_values[key] = value

    next_row = ws.max_row + 1 if ws.max_row and ws.max_row >= 1 else 2
    if ws.max_row < 2:
        next_row = 2

    for col_idx, header_title in enumerate(header_row, start=1):
        value = field_values.get(header_title, "")
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.font = BODY_FONT
        cell.alignment = WRAP_TOP_LEFT

    last_col_letter = get_column_letter(len(header_row))
    ws.auto_filter.ref = f"A1:{last_col_letter}{max(next_row, 2)}"

    wb.save(args.xlsx_path)
    print(f"OK: добавлена строка {next_row} в лист '{args.sheet_name}' файла {args.xlsx_path}.")


if __name__ == "__main__":
    main()
